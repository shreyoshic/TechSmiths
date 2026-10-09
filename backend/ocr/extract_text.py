import os
import json
import hashlib
import csv
import pymupdf
import pytesseract

from PIL import Image, ImageOps, ImageFilter, ImageEnhance
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCUMENTS_DIR = os.path.join(BASE_DIR, "documents")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

os.makedirs(OUTPUT_DIR, exist_ok=True)

SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".log", ".docx",
    ".pptx", ".xlsx", ".csv", ".jpg", ".jpeg", ".png"
}


def calculate_hash(file_path):
    sha256 = hashlib.sha256()

    with open(file_path, "rb") as file:
        for block in iter(lambda: file.read(4096), b""):
            sha256.update(block)

    return sha256.hexdigest()


def fix_encoding(text):
    try:
        return text.encode("latin1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return text


def clean_text(text):
    text = fix_encoding(text)

    return "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )


def make_record(
    filename, reference, text, source_hash,
    method="text", ocr_confidence=None
):
    return {
        "document": filename,
        "page": reference,
        "text": clean_text(text),
        "source_hash": source_hash,
        "extraction_method": method,
        "ocr_confidence": ocr_confidence
    }


def extract_pdf(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)
    records = []

    with pymupdf.open(path) as doc:
        if doc.is_encrypted:
            raise ValueError("PDF is encrypted and cannot be opened.")

        for page_number, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            method = "text"
            confidence = None

            if not text:
                pix = page.get_pixmap(dpi=200, alpha=False)
                image = Image.frombytes(
                    "RGB", (pix.width, pix.height), pix.samples
                )
                image = preprocess_image(image)
                text, confidence = extract_ocr_text(image)

                print(
                f"OCR confidence for {filename}, page {page_number}: "
                f"{confidence}%"
               )

                method = "ocr"

            records.append(
            make_record(
            filename, page_number, text, source_hash,
            method=method,
            ocr_confidence=confidence if method == "ocr" else None
        )
    )

    return records


def extract_text_file(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)

    with open(path, "r", encoding="utf-8-sig", errors="replace") as file:
        text = file.read()

    return [make_record(filename, 1, text, source_hash)]

def preprocess_image(image):
    # Correct image orientation
    image = ImageOps.exif_transpose(image)

    # Convert to grayscale
    image = image.convert("L")

    # Enlarge the image for better text recognition
    image = image.resize(
        (image.width * 2, image.height * 2),
        Image.Resampling.LANCZOS
    )

    # Improve contrast
    image = ImageEnhance.Contrast(image).enhance(2)

    # Reduce small amounts of noise
    image = image.filter(ImageFilter.MedianFilter(size=3))

    return image

def extract_ocr_text(image):
    data = pytesseract.image_to_data(
        image,
        output_type=pytesseract.Output.DICT
    )

    words = []
    confidences = []

    for word, confidence in zip(data["text"], data["conf"]):
        word = word.strip()

        if word and int(confidence) >= 0:
            words.append(word)
            confidences.append(int(confidence))

    text = " ".join(words)

    if confidences:
        average_confidence = sum(confidences) / len(confidences)
    else:
        average_confidence = 0

    return text, round(average_confidence, 2)


def extract_image(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)

    with Image.open(path) as image:
        image = preprocess_image(image)
        text, confidence = extract_ocr_text(image)

        print(f"OCR confidence for {filename}: {confidence}%")

    return [
    make_record(
        filename, 1, text, source_hash,
        method="ocr",
        ocr_confidence=confidence
    )
]


def extract_docx(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)
    doc = Document(path)

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return [make_record(filename, 1, "\n".join(parts), source_hash)]


def extract_pptx(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)
    presentation = Presentation(path)
    records = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)

            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))

        records.append(
            make_record(filename, slide_number, "\n".join(parts), source_hash)
        )

    return records


def extract_xlsx(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []

    try:
        for sheet in workbook.worksheets:
            rows = []

            for row in sheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]

                if any(values):
                    rows.append(" | ".join(values))

            records.append(
                make_record(
                    filename,
                    sheet.title,
                    "\n".join(rows),
                    source_hash
                )
            )
    finally:
        workbook.close()

    return records


def extract_csv(path):
    filename = os.path.basename(path)
    source_hash = calculate_hash(path)

    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as file:
        rows = list(csv.reader(file))

    text = "\n".join(" | ".join(row) for row in rows)
    return [make_record(filename, 1, text, source_hash)]


def extract_file(path):
    extension = os.path.splitext(path)[1].lower()

    if extension == ".pdf":
        return extract_pdf(path)
    if extension in {".txt", ".log"}:
        return extract_text_file(path)
    if extension in {".jpg", ".jpeg", ".png"}:
        return extract_image(path)
    if extension == ".docx":
        return extract_docx(path)
    if extension == ".pptx":
        return extract_pptx(path)
    if extension == ".xlsx":
        return extract_xlsx(path)
    if extension == ".csv":
        return extract_csv(path)

    raise ValueError(f"Unsupported file format: {extension}")


def extract_all_documents():
    all_documents = []

    for filename in sorted(os.listdir(DOCUMENTS_DIR)):
        path = os.path.join(DOCUMENTS_DIR, filename)

        if not os.path.isfile(path):
            continue

        extension = os.path.splitext(filename)[1].lower()

        if extension not in SUPPORTED_EXTENSIONS:
            print(f"Skipped unsupported file: {filename}")
            continue

        try:
            records = extract_file(path)
            all_documents.extend(records)
            print(f"Extracted: {filename} ({len(records)} record(s))")

        except Exception as error:
            print(f"Failed to extract {filename}: {error}")

    combined_path = os.path.join(OUTPUT_DIR, "all_documents.json")

    with open(combined_path, "w", encoding="utf-8") as file:
        json.dump(all_documents, file, indent=4, ensure_ascii=False)

    print(f"Combined output saved to {combined_path}")
    return all_documents


if __name__ == "__main__":
    extract_all_documents()
